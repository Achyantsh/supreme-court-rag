"""Information Retrieval Evaluation Metrics Calculator (V2)

Computes empirical IR metrics (Hit@5, Pooled Recall@5, MRR@5, Graded nDCG@5)
for:
1. Dense Vector Retrieval (bge-small-en-v1.5)
2. Lexical Retrieval (BM25Okapi)
3. Hybrid Retrieval (Reciprocal Rank Fusion, k=60)
4. Multi-stage Hybrid + Cross-Encoder Reranker (ms-marco-MiniLM-L-6-v2)

Methodological Standards:
- Reads evaluated candidates directly from data/eval_candidates_v2.jsonl.
- Does NOT rerun reranking on the fly.
- Reads verified human annotations from data/relevance_review_v2.xlsx by header names.
- Uses stable keys: (question_id, doc_id).
- REJECTS evaluation if any candidate in the evaluated pool or top-5 is unlabelled.
- NEVER converts blank/missing labels into 0.
- Computes ideal DCG from the complete judged pool for each question.
- Recall denominator is the total relevant cases in the judged pool.
- Supports explicit scopes: 'pilot' (Q01-Q05) and 'full' (Q01-Q30).
"""

import argparse
import json
import math
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, Any, List, Tuple, Set

import openpyxl

DEFAULT_CANDIDATES = Path("data/eval_candidates_v2.jsonl")
DEFAULT_WORKBOOK = Path("data/relevance_review_v2.xlsx")
PILOT_METRICS_OUTPUT = Path("data/pilot_metrics_v2.json")
FULL_METRICS_OUTPUT = Path("data/full_benchmark_metrics_v2.json")

PILOT_QIDS = [f"Q{i:02d}" for i in range(1, 6)]
FULL_QIDS = [f"Q{i:02d}" for i in range(1, 31)]
METHODS = ["dense", "bm25", "hybrid_rrf", "hybrid_reranked"]


def dcg_at_k(rel_scores: List[int], k: int = 5) -> float:
    """Compute Discounted Cumulative Gain at rank K using standard exponential gain."""
    return sum(
        (2**rel - 1) / math.log2(idx + 1)
        for idx, rel in enumerate(rel_scores[:k], start=1)
    )


def idcg_at_k(pool_scores: List[int], k: int = 5) -> float:
    """Compute Ideal Discounted Cumulative Gain at rank K from the sorted judged pool."""
    sorted_ideal = sorted(pool_scores, reverse=True)
    return dcg_at_k(sorted_ideal, k=k)


def ndcg_at_k(rel_scores: List[int], pool_scores: List[int], k: int = 5) -> float:
    """Compute Normalized Discounted Cumulative Gain at rank K.
    
    If no relevant documents exist in the judged pool (IDCG == 0):
      Convention: returns 1.0 if retrieved DCG is 0 (optimal choice), else 0.0.
    """
    dcg = dcg_at_k(rel_scores, k=k)
    idcg = idcg_at_k(pool_scores, k=k)
    if idcg > 0:
        return dcg / idcg
    return 1.0 if dcg == 0.0 else 0.0


def load_ground_truth(
    workbook_path: Path,
    expected_qids: Set[str],
) -> Tuple[Dict[Tuple[str, str], int], Dict[str, List[int]], List[Dict[str, Any]]]:
    """Load ground truth labels from the canonical review sheet by header names.
    
    Returns:
      labels: (qid, doc_id) -> int relevance (0, 1, 2)
      pool_by_qid: qid -> list of all judged relevance scores in the pool
      unlabelled_candidates: list of records with missing/blank labels for expected_qids
    """
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing review workbook: {workbook_path}")

    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet_name = "Relevance Review"
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found in {workbook_path}")

    ws = wb[sheet_name]
    header_row = [c for c in next(ws.iter_rows(values_only=True), [])]

    def get_col(name):
        try:
            return header_row.index(name)
        except ValueError:
            raise ValueError(f"Missing required header column '{name}' in {sheet_name}")

    qid_col = get_col("Question ID")
    doc_id_col = get_col("doc_id")
    rel_col = get_col("Relevance (0/1/2)")
    case_col = get_col("Case name")
    cit_col = get_col("Citation")

    labels = {}
    pool_by_qid = defaultdict(list)
    unlabelled = []

    for row_idx, row in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2):
        qid = str(row[qid_col] or "").strip()
        if qid not in expected_qids:
            continue

        doc_id = str(row[doc_id_col] or "").strip()
        rel_val = row[rel_col]
        case_name = str(row[case_col] or "").strip()
        citation = str(row[cit_col] or "").strip()

        if rel_val is None or str(rel_val).strip() == "":
            unlabelled.append({
                "row_idx": row_idx,
                "question_id": qid,
                "doc_id": doc_id,
                "case_name": case_name,
                "citation": citation,
            })
        else:
            try:
                numeric_rel = int(rel_val)
                if numeric_rel not in (0, 1, 2):
                    unlabelled.append({
                        "row_idx": row_idx,
                        "question_id": qid,
                        "doc_id": doc_id,
                        "case_name": case_name,
                        "citation": citation,
                        "error": f"Invalid value {rel_val}; must be 0, 1, or 2",
                    })
                else:
                    labels[(qid, doc_id)] = numeric_rel
                    pool_by_qid[qid].append(numeric_rel)
            except ValueError:
                unlabelled.append({
                    "row_idx": row_idx,
                    "question_id": qid,
                    "doc_id": doc_id,
                    "case_name": case_name,
                    "citation": citation,
                    "error": f"Non-integer value {rel_val}",
                })

    wb.close()
    return labels, pool_by_qid, unlabelled


def load_candidate_rankings(
    candidates_file: Path,
    expected_qids: Set[str],
) -> Dict[str, Dict[str, List[Dict[str, Any]]]]:
    """Load persisted candidate rankings from eval_candidates_v2.jsonl.
    
    Returns:
      data[qid][method] = list of ranked result dicts
    """
    if not candidates_file.exists():
        raise FileNotFoundError(f"Missing candidates file: {candidates_file}")

    data = defaultdict(dict)
    with candidates_file.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            qid = row["question_id"]
            if qid in expected_qids:
                method = row["method"]
                data[qid][method] = row["results"]

    # Validate that all expected questions have all methods
    missing_records = []
    for qid in sorted(expected_qids):
        for m in METHODS:
            if m not in data[qid]:
                missing_records.append(f"{qid}:{m}")

    if missing_records:
        raise ValueError(
            f"Candidates file is missing ranking records for: {', '.join(missing_records)}"
        )

    return data


def verify_completeness_and_reject(
    expected_qids: List[str],
    retrieval_data: Dict[str, Dict[str, List[Dict[str, Any]]]],
    labels: Dict[Tuple[str, str], int],
    unlabelled_in_workbook: List[Dict[str, Any]],
    k: int = 5,
):
    """Enforce complete evaluation labels. Rejects immediately if ANY top-K or pool label is missing."""
    missing_top_k = []

    for qid in expected_qids:
        for m in METHODS:
            results = retrieval_data[qid][m][:k]
            for r in results:
                doc_id = r["doc_id"]
                key = (qid, doc_id)
                if key not in labels:
                    missing_top_k.append({
                        "question_id": qid,
                        "doc_id": doc_id,
                        "case_name": r.get("case_name"),
                        "method": m,
                        "rank": r.get("rank"),
                    })

    if unlabelled_in_workbook or missing_top_k:
        print("\n" + "=" * 80)
        print("CRITICAL VALIDATION ERROR: INCOMPLETE HUMAN RELEVANCE LABELS")
        print("=" * 80)
        print("In accordance with scientific evaluation standards:")
        print("1. Missing/blank labels are NEVER silently assigned relevance 0.")
        print("2. Unanswered evaluation questions are NEVER silently excluded.")
        print("3. Final benchmark metrics cannot be written until all required candidates are reviewed.\n")

        if unlabelled_in_workbook:
            print(f"Found {len(unlabelled_in_workbook)} unlabelled candidate(s) in review sheet:")
            for item in unlabelled_in_workbook[:10]:
                print(f"  - Row {item['row_idx']} | {item['question_id']} | {item['doc_id']} | {item['case_name']}")
            if len(unlabelled_in_workbook) > 10:
                print(f"  ... and {len(unlabelled_in_workbook) - 10} more.")

        if missing_top_k:
            print(f"\nFound {len(missing_top_k)} top-{k} retrieved candidates missing from labels index:")
            for item in missing_top_k[:10]:
                print(f"  - {item['question_id']} | {item['method']} rank {item['rank']} | {item['doc_id']} | {item['case_name']}")
            if len(missing_top_k) > 10:
                print(f"  ... and {len(missing_top_k) - 10} more.")

        print("\nACTION REQUIRED:")
        print("Open 'data/relevance_review_v2.xlsx', complete the missing relevance labels (0, 1, or 2), and re-run.")
        print("=" * 80)
        raise RuntimeError(
            f"Evaluation halted: {len(unlabelled_in_workbook)} unlabelled pool rows and "
            f"{len(missing_top_k)} missing top-{k} labels detected."
        )


def evaluate_scope(
    scope: str = "pilot",
    k: int = 5,
    candidates_file: Path = DEFAULT_CANDIDATES,
    workbook_file: Path = DEFAULT_WORKBOOK,
    output_file: Path = None,
) -> Dict[str, Any]:
    """Compute and macro-average IR evaluation metrics for the specified scope."""
    if scope == "pilot":
        expected_qids = PILOT_QIDS
        default_out = PILOT_METRICS_OUTPUT
    elif scope == "full":
        expected_qids = FULL_QIDS
        default_out = FULL_METRICS_OUTPUT
    else:
        raise ValueError(f"Unknown scope: '{scope}'. Use 'pilot' or 'full'.")

    target_output = output_file or default_out
    qids_set = set(expected_qids)

    print("=" * 85)
    print(f"EVALUATION METRICS CALCULATOR (V2) | Scope: {scope.upper()} ({len(expected_qids)} Questions, K={k})")
    print("=" * 85)

    # 1. Load candidate rankings
    retrieval_data = load_candidate_rankings(candidates_file, qids_set)
    print(f"Loaded rankings for {len(retrieval_data)} questions across {len(METHODS)} methods.")

    # 2. Load ground truth labels
    labels, pool_by_qid, unlabelled = load_ground_truth(workbook_file, qids_set)
    print(f"Loaded {len(labels)} human relevance labels from {workbook_file.name}.")

    # 3. Reject incomplete annotations
    verify_completeness_and_reject(
        expected_qids=expected_qids,
        retrieval_data=retrieval_data,
        labels=labels,
        unlabelled_in_workbook=unlabelled,
        k=k,
    )

    # 4. Compute metrics
    metrics_by_method = {
        m: {
            "hit_lenient": [],
            "hit_strict": [],
            "recall_lenient": [],
            "recall_strict": [],
            "mrr_lenient": [],
            "mrr_strict": [],
            "ndcg": [],
            "per_question": {},
        }
        for m in METHODS
    }

    zero_relevance_queries = {
        "lenient": [],
        "strict": [],
    }

    for qid in expected_qids:
        pool_ratings = pool_by_qid[qid]
        tot_lenient = sum(1 for r in pool_ratings if r >= 1)
        tot_strict = sum(1 for r in pool_ratings if r >= 2)

        if tot_lenient == 0 and qid not in zero_relevance_queries["lenient"]:
            zero_relevance_queries["lenient"].append(qid)
        if tot_strict == 0 and qid not in zero_relevance_queries["strict"]:
            zero_relevance_queries["strict"].append(qid)

        for m in METHODS:
            results = retrieval_data[qid][m][:k]
            rel_scores = [labels[(qid, r["doc_id"])] for r in results]

            # Hit Rate@K
            h_lenient = 1.0 if any(s >= 1 for s in rel_scores) else 0.0
            h_strict = 1.0 if any(s >= 2 for s in rel_scores) else 0.0

            # Pooled Recall@K (denominator: all relevant items in the judged pool)
            ret_lenient = sum(1 for s in rel_scores if s >= 1)
            ret_strict = sum(1 for s in rel_scores if s >= 2)

            r_lenient = (ret_lenient / tot_lenient) if tot_lenient > 0 else 0.0
            r_strict = (ret_strict / tot_strict) if tot_strict > 0 else 0.0

            # MRR@K
            mrr_lenient = 0.0
            for rank_idx, s in enumerate(rel_scores, start=1):
                if s >= 1:
                    mrr_lenient = 1.0 / rank_idx
                    break

            mrr_strict = 0.0
            for rank_idx, s in enumerate(rel_scores, start=1):
                if s >= 2:
                    mrr_strict = 1.0 / rank_idx
                    break

            # Graded nDCG@K using complete judged pool for IDCG
            ndcg = ndcg_at_k(rel_scores=rel_scores, pool_scores=pool_ratings, k=k)

            metrics_by_method[m]["hit_lenient"].append(h_lenient)
            metrics_by_method[m]["hit_strict"].append(h_strict)
            metrics_by_method[m]["recall_lenient"].append(r_lenient)
            metrics_by_method[m]["recall_strict"].append(r_strict)
            metrics_by_method[m]["mrr_lenient"].append(mrr_lenient)
            metrics_by_method[m]["mrr_strict"].append(mrr_strict)
            metrics_by_method[m]["ndcg"].append(ndcg)

            metrics_by_method[m]["per_question"][qid] = {
                "rel_scores": rel_scores,
                "hit_strict": h_strict,
                "hit_lenient": h_lenient,
                "recall_strict": round(r_strict, 4),
                "recall_lenient": round(r_lenient, 4),
                "mrr_strict": round(mrr_strict, 4),
                "mrr_lenient": round(mrr_lenient, 4),
                "ndcg": round(ndcg, 4),
            }

    # Macro averages
    num_q = len(expected_qids)
    summary = {}

    print("\n" + "=" * 90)
    print(f"BENCHMARK RESULTS (Macro-Averaged across {num_q} Questions, K={k})")
    print("=" * 90)
    print(
        f"{'Retrieval Method':<18} | {'Hit@5(L)':<8} | {'Rec@5(L)':<9} | {'MRR@5(L)':<8} | "
        f"{'nDCG@5':<8} | {'Strict Hit':<10} | {'Strict MRR':<10} | {'Strict Rec':<10}"
    )
    print("-" * 90)

    for m in METHODS:
        data = metrics_by_method[m]
        avg_hl = sum(data["hit_lenient"]) / num_q
        avg_hs = sum(data["hit_strict"]) / num_q
        avg_rl = sum(data["recall_lenient"]) / num_q
        avg_rs = sum(data["recall_strict"]) / num_q
        avg_mrrl = sum(data["mrr_lenient"]) / num_q
        avg_mrrs = sum(data["mrr_strict"]) / num_q
        avg_ndcg = sum(data["ndcg"]) / num_q

        summary[m] = {
            "hit_rate_lenient": round(avg_hl, 4),
            "hit_rate_strict": round(avg_hs, 4),
            "recall_lenient": round(avg_rl, 4),
            "recall_strict": round(avg_rs, 4),
            "mrr_lenient": round(avg_mrrl, 4),
            "mrr_strict": round(avg_mrrs, 4),
            "ndcg_at_5": round(avg_ndcg, 4),
            "per_question": data["per_question"],
        }

        print(
            f"{m:<18} | {avg_hl:<8.4f} | {avg_rl:<9.4f} | {avg_mrrl:<8.4f} | "
            f"{avg_ndcg:<8.4f} | {avg_hs:<10.4f} | {avg_mrrs:<10.4f} | {avg_rs:<10.4f}"
        )

    print("=" * 90)
    print("Scientific Methodology Notes:")
    print("  * (L) = Lenient threshold (Relevance >= 1: Partial or Strong).")
    print("  * Strict = Strict threshold (Relevance >= 2: Strong only).")
    print("  * Graded nDCG@5 uses gain 2^rel - 1 and log2(rank + 1) discount against complete pool IDCG.")
    print("  * LIMITATION: Recall@5 is pooled recall calculated against the judged pool (union of top-5),")
    print("    not against unretrieved judgments across the broader 800-case corpus.")
    if zero_relevance_queries["strict"]:
        print(f"  * Strict zero-relevance queries (tot_strict=0 in pool): {zero_relevance_queries['strict']}")
    print("=" * 90)

    report = {
        "benchmark_scope": scope,
        "k": k,
        "questions_count": num_q,
        "questions": expected_qids,
        "methodology_disclaimer": (
            "Pooled Recall@5 is evaluated against the reviewed judged pool (union of top-5 candidates), "
            "not the complete unreviewed 800-judgment corpus."
        ),
        "zero_relevance_queries": zero_relevance_queries,
        "summary": summary,
    }

    target_output.parent.mkdir(parents=True, exist_ok=True)
    with target_output.open("w", encoding="utf-8") as f:
        json.dump(report, f, indent=2)

    print(f"\nMetrics written to: {target_output.resolve()}")
    return report


def main():
    parser = argparse.ArgumentParser(description="Supreme Court RAG Retrieval Metrics Calculator")
    parser.add_argument(
        "--scope",
        choices=["pilot", "full"],
        default="pilot",
        help="Evaluation scope: 'pilot' (Q01-Q05) or 'full' (Q01-Q30)",
    )
    parser.add_argument(
        "-k",
        "--k",
        type=int,
        default=5,
        help="Cutoff rank K (default: 5)",
    )
    parser.add_argument(
        "--candidates",
        type=Path,
        default=DEFAULT_CANDIDATES,
        help="Path to evaluation candidates JSONL",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to relevance review workbook",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Custom output JSON path",
    )
    args = parser.parse_args()

    try:
        evaluate_scope(
            scope=args.scope,
            k=args.k,
            candidates_file=args.candidates,
            workbook_file=args.workbook,
            output_file=args.output,
        )
    except Exception as e:
        print(f"\n[EXECUTION HALTED] {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
