"""Formatted Excel Relevance Review Workbook Generator

Builds a single canonical review sheet from the 4 retrieval methods:
1. Dense Vector (bge-small-en-v1.5)
2. Lexical (BM25Okapi)
3. Hybrid RRF (k=60)
4. Multi-stage Hybrid + Cross-Encoder Reranker (ms-marco-MiniLM-L-6-v2)

Features:
- Takes the union of top-5 candidates from all 4 methods.
- Deduplicates by (Question ID, doc_id).
- Preserves individual method ranks and tagged excerpts.
- Preserves any user-entered annotations when re-run.
- Migrates existing pilot human labels from data/relevance_review_pilot_labeled.xlsx.
- Logs any ambiguous or unmapped legacy labels to an explicit audit sheet.
- Validates dropdown entries (0, 1, 2) and conditional color formatting.
"""

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, Any, List, Tuple

import openpyxl
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INPUT_CANDIDATES = Path("data/eval_candidates_v2.jsonl")
LEGACY_WORKBOOK = Path("data/relevance_review_pilot_labeled.xlsx")
OUTPUT_WORKBOOK = Path("data/relevance_review_v2.xlsx")

METHODS = ["dense", "bm25", "hybrid_rrf", "hybrid_reranked"]
PILOT_QUESTIONS = {"Q01", "Q02", "Q03", "Q04", "Q05"}
REVIEW_RANK_CUTOFF = 5

COLUMNS = [
    "Question ID",
    "Topic",
    "Question",
    "doc_id",
    "Case name",
    "Citation",
    "Decision date",
    "Dense rank",
    "BM25 rank",
    "Hybrid RRF rank",
    "Hybrid reranked rank",
    "Evidence excerpts",
    "Relevance (0/1/2)",
    "Review notes",
]


def normalize_text(text: str) -> str:
    """Normalize text for safe string matching, standardizing legal abbreviations like v./versus."""
    if not text:
        return ""
    cleaned = re.sub(r"[^a-z0-9]+", " ", str(text).lower())
    words = cleaned.split()
    # Standardize variations like versus / vs / v
    standardized = ["v" if w in ("versus", "vs") else w for w in words]
    return " ".join(standardized)


def clean_snippet(text: str) -> str:
    """Clean whitespace in passage excerpts."""
    return re.sub(r"\s+", " ", text or "").strip()


def question_number(qid: str) -> int:
    match = re.search(r"\d+", qid or "")
    return int(match.group()) if match else 9999


def load_existing_v2_labels(workbook_path: Path = OUTPUT_WORKBOOK) -> Dict[Tuple[str, str], Tuple[Any, str]]:
    """Load labels already entered in an existing v2 workbook to prevent accidental overwrite."""
    if not workbook_path.exists():
        return {}

    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=True)
    except Exception as e:
        print(f"Notice: Could not load existing v2 workbook ({e}). Proceeding afresh.")
        return {}

    labels = {}
    sheet_name = "Relevance Review"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [c for c in next(ws.iter_rows(values_only=True), [])]
        try:
            qid_idx = headers.index("Question ID")
            doc_idx = headers.index("doc_id")
            rel_idx = headers.index("Relevance (0/1/2)")
            notes_idx = headers.index("Review notes")

            for row in list(ws.iter_rows(values_only=True))[1:]:
                qid = str(row[qid_idx] or "").strip()
                doc_id = str(row[doc_idx] or "").strip()
                rel = row[rel_idx]
                notes = str(row[notes_idx] or "").strip()
                if qid and doc_id and rel is not None and str(rel).strip() != "":
                    try:
                        labels[(qid, doc_id)] = (int(rel), notes)
                    except ValueError:
                        pass
        except ValueError:
            pass

    wb.close()
    return labels


def load_legacy_annotations(legacy_path: Path = LEGACY_WORKBOOK) -> List[Dict[str, Any]]:
    """Load annotations from the legacy pilot review workbook."""
    if not legacy_path.exists():
        return []

    wb = openpyxl.load_workbook(legacy_path, data_only=True)
    if "Pilot Review" not in wb.sheetnames:
        wb.close()
        return []

    ws = wb["Pilot Review"]
    headers = [c for c in next(ws.iter_rows(values_only=True), [])]
    try:
        qid_col = headers.index("Question ID")
        case_col = headers.index("Case name")
        cit_col = headers.index("Citation")
        rel_col = headers.index("Relevance (0/1/2)")
        notes_col = headers.index("Review notes")
    except ValueError as e:
        print(f"Warning: Unexpected columns in legacy workbook: {e}")
        wb.close()
        return []

    annotations = []
    for row in list(ws.iter_rows(values_only=True))[1:]:
        qid = str(row[qid_col] or "").strip()
        case_name = str(row[case_col] or "").strip()
        citation = str(row[cit_col] or "").strip()
        rel = row[rel_col]
        notes = str(row[notes_col] or "").strip()

        if qid and rel is not None and str(rel).strip() != "":
            try:
                rel_val = int(rel)
                annotations.append({
                    "question_id": qid,
                    "case_name": case_name,
                    "citation": citation,
                    "relevance": rel_val,
                    "notes": notes,
                })
            except ValueError:
                pass

    wb.close()
    return annotations


def pool_candidates(
    candidates_file: Path = INPUT_CANDIDATES,
    rank_cutoff: int = REVIEW_RANK_CUTOFF,
) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """Pool top-K candidates from all 4 methods and deduplicate by (question_id, doc_id)."""
    if not candidates_file.exists():
        raise FileNotFoundError(
            f"Missing candidates file: {candidates_file}. Run 'python eval.py' first."
        )

    pooled = {}
    with candidates_file.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            qid = row["question_id"]
            method = row["method"]
            topic = row.get("topic", "")
            question = row.get("question", "")

            for item in row.get("results", []):
                rank = int(item["rank"])
                if rank > rank_cutoff:
                    continue

                doc_id = str(item.get("doc_id") or item.get("case_name") or "unknown").strip()
                key = (qid, doc_id)

                if key not in pooled:
                    pooled[key] = {
                        "Question ID": qid,
                        "Topic": topic,
                        "Question": question,
                        "doc_id": doc_id,
                        "Case name": item.get("case_name") or "",
                        "Citation": item.get("citation") or "",
                        "Decision date": item.get("decision_date") or "",
                        "Dense rank": "",
                        "BM25 rank": "",
                        "Hybrid RRF rank": "",
                        "Hybrid reranked rank": "",
                        "excerpts_list": [],
                        "Relevance (0/1/2)": "",
                        "Review notes": "",
                    }

                record = pooled[key]
                rank_col_map = {
                    "dense": "Dense rank",
                    "bm25": "BM25 rank",
                    "hybrid_rrf": "Hybrid RRF rank",
                    "hybrid_reranked": "Hybrid reranked rank",
                }
                if method in rank_col_map:
                    record[rank_col_map[method]] = rank

                snippet = clean_snippet(item.get("text", ""))
                if snippet:
                    tagged = f"[{method}, rank {rank}] {snippet}"
                    if tagged not in record["excerpts_list"]:
                        record["excerpts_list"].append(tagged)

    # Format excerpts (keep up to 2 distinct excerpts)
    for record in pooled.values():
        record["Evidence excerpts"] = "\n\n".join(record["excerpts_list"][:2])

    return pooled


def migrate_and_merge_labels(
    pooled: Dict[Tuple[str, str], Dict[str, Any]],
    legacy_annotations: List[Dict[str, Any]],
    existing_v2_labels: Dict[Tuple[str, str], Tuple[Any, str]],
) -> Tuple[int, int, List[Dict[str, Any]], int]:
    """Migrate legacy annotations and preserve any existing v2 user edits."""
    migrated_count = 0
    unresolved_legacy = []

    # 1. First apply existing v2 user edits (highest priority)
    for key, (rel, notes) in existing_v2_labels.items():
        if key in pooled:
            pooled[key]["Relevance (0/1/2)"] = rel
            if notes and not pooled[key]["Review notes"]:
                pooled[key]["Review notes"] = notes

    # Group pooled candidates by question_id for index lookup
    candidates_by_qid = defaultdict(list)
    for key, cand in pooled.items():
        candidates_by_qid[cand["Question ID"]].append(cand)

    # 2. Migrate legacy annotations
    for leg in legacy_annotations:
        qid = leg["question_id"]
        leg_case = normalize_text(leg["case_name"])
        leg_cit = normalize_text(leg["citation"])
        rel = leg["relevance"]
        notes = leg["notes"]

        matching_candidates = candidates_by_qid.get(qid, [])

        # Step 2a: Exact normalized (case_name, citation) match
        exact_matches = [
            c for c in matching_candidates
            if normalize_text(c["Case name"]) == leg_case
            and normalize_text(c["Citation"]) == leg_cit
        ]

        target_candidate = None
        if len(exact_matches) == 1:
            target_candidate = exact_matches[0]
        elif len(exact_matches) > 1:
            unresolved_legacy.append({
                **leg,
                "reason": f"Ambiguous match: {len(exact_matches)} candidates match (case, citation)",
            })
            continue
        else:
            # Step 2b: Fallback to case_name-only ONLY IF uniquely matched
            case_matches = [
                c for c in matching_candidates
                if normalize_text(c["Case name"]) == leg_case
            ]
            if len(case_matches) == 1:
                target_candidate = case_matches[0]
            elif len(case_matches) > 1:
                unresolved_legacy.append({
                    **leg,
                    "reason": f"Ambiguous fallback: {len(case_matches)} candidates match case_name",
                })
                continue
            else:
                unresolved_legacy.append({
                    **leg,
                    "reason": "Not in pooled candidates (not retrieved in Top-5 by any of the 4 methods)",
                })
                continue

        if target_candidate is not None:
            # If not already labeled by v2 existing edits, populate
            if target_candidate["Relevance (0/1/2)"] == "":
                target_candidate["Relevance (0/1/2)"] = rel
                if notes and not target_candidate["Review notes"]:
                    target_candidate["Review notes"] = notes
            migrated_count += 1

    # Count unlabelled
    unlabelled_total = sum(1 for c in pooled.values() if c["Relevance (0/1/2)"] == "")
    unlabelled_pilot = sum(
        1 for c in pooled.values()
        if c["Question ID"] in PILOT_QUESTIONS and c["Relevance (0/1/2)"] == ""
    )

    return migrated_count, len(unresolved_legacy), unresolved_legacy, unlabelled_pilot


def format_workbook(wb: Workbook, pooled_items: List[Dict[str, Any]], unresolved: List[Dict[str, Any]]):
    """Format and style Excel review sheets."""
    # Sheet 1: Instructions
    instructions = wb.active
    instructions.title = "Instructions"

    guide_lines = [
        ["SUPREME COURT RAG — RELEVANCE REVIEW WORKBOOK (V2)"],
        [""],
        ["WORKFLOW"],
        ["1. Open the 'Relevance Review' sheet."],
        ["2. Filter 'Question ID' for the Pilot set (Q01 to Q05) first."],
        ["3. Verify existing pilot annotations and label any newly retrieved candidates."],
        ["4. Run 'python calculate_metrics.py --scope pilot' once the pilot is complete."],
        ["5. Proceed to remaining questions (Q06 to Q30) for full benchmark evaluation."],
        [""],
        ["LABEL DEFINITIONS"],
        ["0 = Irrelevant: Does not meaningfully address or answer the legal query."],
        ["1 = Partial: Related legal context or principles, but incomplete/indirect."],
        ["2 = Strong: Directly, substantively, and clearly addresses the core legal inquiry."],
        [""],
        ["RULES FOR EVALUATION RIGOR"],
        ["- Do NOT assign higher scores merely because a case appears in multiple methods or has rank 1."],
        ["- Every candidate appears exactly once with all its retrieval method ranks indicated."],
        ["- Never leave labels blank when completing a benchmark scope; missing labels reject evaluation."],
        ["- Rerunning 'python prepare_labels.py' will safely preserve all entered labels."],
    ]
    for row in guide_lines:
        instructions.append(row)

    instructions.column_dimensions["A"].width = 95
    instructions.sheet_view.showGridLines = False
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    instructions["A1"].font = Font(bold=True, size=15, color="1F497D")
    for r_idx in [3, 10, 16]:
        instructions.cell(row=r_idx, column=1).font = Font(bold=True, size=12, color="243447")

    # Sheet 2: Canonical Relevance Review
    ws = wb.create_sheet("Relevance Review")
    ws.append(COLUMNS)

    for item in pooled_items:
        ws.append([
            item["Question ID"],
            item["Topic"],
            item["Question"],
            item["doc_id"],
            item["Case name"],
            item["Citation"],
            item["Decision date"],
            item["Dense rank"],
            item["BM25 rank"],
            item["Hybrid RRF rank"],
            item["Hybrid reranked rank"],
            item["Evidence excerpts"],
            item["Relevance (0/1/2)"],
            item["Review notes"],
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = True

    header_fill = PatternFill("solid", fgColor="1F497D")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    col_widths = {
        "A": 12,  # Question ID
        "B": 22,  # Topic
        "C": 45,  # Question
        "D": 22,  # doc_id
        "E": 38,  # Case name
        "F": 22,  # Citation
        "G": 14,  # Decision date
        "H": 12,  # Dense rank
        "I": 12,  # BM25 rank
        "J": 14,  # Hybrid RRF rank
        "K": 16,  # Hybrid reranked rank
        "L": 85,  # Evidence excerpts
        "M": 18,  # Relevance (0/1/2)
        "N": 35,  # Review notes
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 80
        for cell in ws[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    rel_col_idx = COLUMNS.index("Relevance (0/1/2)") + 1
    rel_letter = get_column_letter(rel_col_idx)

    # Validation dropdown (0, 1, 2)
    validation = DataValidation(
        type="list",
        formula1='"0,1,2"',
        allow_blank=True,
    )
    validation.error = "Select 0 (Irrelevant), 1 (Partial), or 2 (Strong)."
    validation.errorTitle = "Invalid Relevance Rating"
    ws.add_data_validation(validation)
    validation.add(f"{rel_letter}2:{rel_letter}{ws.max_row}")

    # Color formatting
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['"0"'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['0'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['"1"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['1'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['"2"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    ws.conditional_formatting.add(
        f"{rel_letter}2:{rel_letter}{ws.max_row}",
        CellIsRule(operator="equal", formula=['2'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )

    # Sheet 3: Unmapped Legacy Labels (if any)
    if unresolved:
        ws_unmapped = wb.create_sheet("Unmapped Legacy Labels")
        unmapped_cols = ["Question ID", "Case name", "Citation", "Relevance", "Review notes", "Reason"]
        ws_unmapped.append(unmapped_cols)
        for item in unresolved:
            ws_unmapped.append([
                item.get("question_id"),
                item.get("case_name"),
                item.get("citation"),
                item.get("relevance"),
                item.get("notes"),
                item.get("reason"),
            ])
        ws_unmapped.freeze_panes = "A2"
        for cell in ws_unmapped[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C00000")
        for col_letter, w in zip(["A", "B", "C", "D", "E", "F"], [12, 35, 20, 10, 30, 45]):
            ws_unmapped.column_dimensions[col_letter].width = w


def generate_review_workbook(
    candidates_file: Path = INPUT_CANDIDATES,
    legacy_workbook: Path = LEGACY_WORKBOOK,
    output_workbook: Path = OUTPUT_WORKBOOK,
):
    print("=" * 75)
    print("RELEVANCE REVIEW WORKBOOK BUILDER (V2)")
    print("=" * 75)

    pooled_map = pool_candidates(candidates_file=candidates_file)
    print(f"Total unique (Question ID, doc_id) pooled candidates: {len(pooled_map)}")

    existing_v2 = load_existing_v2_labels(workbook_path=output_workbook)
    if existing_v2:
        print(f"Loaded {len(existing_v2)} previously entered labels from existing v2 workbook.")

    legacy_annotations = load_legacy_annotations(legacy_path=legacy_workbook)
    print(f"Found {len(legacy_annotations)} legacy annotations in {legacy_workbook.name}.")

    migrated, unresolved_cnt, unresolved_list, unlabelled_pilot = migrate_and_merge_labels(
        pooled=pooled_map,
        legacy_annotations=legacy_annotations,
        existing_v2_labels=existing_v2,
    )

    sorted_items = sorted(
        pooled_map.values(),
        key=lambda item: (
            question_number(item["Question ID"]),
            item["Case name"],
        ),
    )

    wb = Workbook()
    format_workbook(wb, sorted_items, unresolved_list)

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output_workbook)
    except PermissionError:
        print(f"\n[ERROR] Permission denied: '{output_workbook}'.")
        print("Please close Microsoft Excel and re-run.")
        raise

    total_unlabelled = sum(1 for c in sorted_items if c["Relevance (0/1/2)"] == "")
    total_labelled = len(sorted_items) - total_unlabelled

    print("\n--- MIGRATION & POOLING REPORT ---")
    print(f"Legacy annotations found:        {len(legacy_annotations)}")
    print(f"Legacy annotations migrated:     {migrated}")
    print(f"Legacy annotations unresolved:   {unresolved_cnt}")
    print(f"Total pooled candidates (V2):    {len(sorted_items)}")
    print(f"Total labelled candidates:       {total_labelled}")
    print(f"Total unlabelled candidates:     {total_unlabelled}")
    print(f"Unlabelled in Pilot (Q01-Q05):   {unlabelled_pilot}")
    print(f"Workbook generated successfully: {output_workbook.resolve()}")

    if unresolved_cnt > 0:
        print(f"Notice: {unresolved_cnt} legacy item(s) logged to 'Unmapped Legacy Labels' sheet.")

    return sorted_items


def main():
    generate_review_workbook()


if __name__ == "__main__":
    main()
